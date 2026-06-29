"""Excel QA 報表產生器。

輸出 ``report.xlsx``：每列一個建案，每欄一個比對欄位，儲存格以顏色標示
PASS / WARNING / FAIL，並透過註解附上雙方數值與說明。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.compare.fields import FIELD_SPECS
from app.models.diff_result import QAReport, Status
from app.reports.base import Reporter
from app.utils.logger import get_logger

logger = get_logger(__name__)

_FILLS = {
    Status.PASS: PatternFill("solid", fgColor="C8E6C9"),
    Status.WARNING: PatternFill("solid", fgColor="FFF2CC"),
    Status.FAIL: PatternFill("solid", fgColor="F8C9C9"),
}
_HEADER_FILL = PatternFill("solid", fgColor="1F2329")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


class ExcelReporter(Reporter):
    """產生 ``report.xlsx``。"""

    name = "excel"
    extension = "xlsx"

    def generate(self, report: QAReport, output_dir: Path) -> Path:
        """產生 Excel 報表並回傳檔案路徑。"""
        output_dir.mkdir(parents=True, exist_ok=True)
        path = output_dir / "report.xlsx"

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "QA 比對"

        headers = ["建案", "編號", "整體"] + [spec.label for spec in FIELD_SPECS]
        sheet.append(headers)
        for col in range(1, len(headers) + 1):
            cell = sheet.cell(row=1, column=col)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        for project in report.projects:
            field_by_key = {f.key: f for f in project.fields}
            row = [project.name, project.project_id, project.status.value]
            row += [
                field_by_key[spec.key].status.value if spec.key in field_by_key else ""
                for spec in FIELD_SPECS
            ]
            sheet.append(row)
            self._style_row(sheet, sheet.max_row, project, field_by_key)

        self._autosize(sheet, headers)
        sheet.freeze_panes = "D2"
        workbook.save(path)
        logger.info("已輸出 Excel 報表：%s", path)
        return path

    def _style_row(self, sheet, row_idx, project, field_by_key) -> None:
        """為單列套用狀態顏色與註解（註解列出每個來源的值與說明）。"""
        sheet.cell(row=row_idx, column=3).fill = _FILLS.get(
            project.status, _FILLS[Status.WARNING]
        )
        for offset, spec in enumerate(FIELD_SPECS):
            field = field_by_key.get(spec.key)
            if not field:
                continue
            cell = sheet.cell(row=row_idx, column=4 + offset)
            cell.fill = _FILLS[field.status]
            cell.alignment = Alignment(horizontal="center")
            lines = [f"HouseQA：{field.house_value}"]
            for comp in field.comparisons:
                lines.append(f"{comp.source}（{comp.status.value}）：{comp.value}｜{comp.message}")
            cell.comment = Comment("\n".join(lines), "HouseQA")

    @staticmethod
    def _autosize(sheet, headers) -> None:
        """簡易欄寬設定。"""
        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 10
        for col in range(3, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 10
